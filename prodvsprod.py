#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SFDC Production Metadata Dump Comparison – end-to-end builder

Fixes:
- Latest dump is selected by NAME date token (not mtime).
- ConfigParser interpolation disabled (handles %20 safely).
- XLSX-only outputs (no CSV anywhere). Hard fail if cannot write XLSX.
- SharePoint link forced to .xlsx and "(Execution Date)" removed.
- codecomp.html beautified to a minimal, single-table page (header, table, footer).
- Graceful, visible no-op when no new dump (email + marker + exit 7).
"""

import os
import re
import sys
import json
import shutil
import hashlib
import subprocess
import datetime as dt
from collections import Counter
from configparser import ConfigParser
from pathlib import Path
from email.mime.text import MIMEText
import smtplib

# =========================
# Basic file & path helpers
# =========================
def load_conf(p: Path) -> ConfigParser:
    if not p.exists():
        sys.exit(f"[ERROR] Missing config: {p}")
    # Disable interpolation so %20 etc. in URLs are read literally
    c = ConfigParser(interpolation=None)
    c.read(p)
    return c

def read_text(p: Path) -> str:
    return p.read_text(encoding="utf-8") if p.exists() else ""

def write_text(p: Path, s: str):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(s, encoding="utf-8")

def append_text(p: Path, s: str):
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("a", encoding="utf-8") as f:
        f.write(s)

def copy_tree(src: Path, dst: Path):
    if dst.exists():
        shutil.rmtree(dst)
    shutil.copytree(src, dst)

def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for ch in iter(lambda: f.read(1024 * 1024), b""):
            h.update(ch)
    return h.hexdigest()

def walk_files(base: Path):
    for r, _, fs in os.walk(base):
        for f in fs:
            full = Path(r) / f
            yield full.relative_to(base), full

def top_level_dirs(p: Path) -> set:
    return {x.name for x in p.iterdir() if x.is_dir()}

def today_iso() -> str:
    return dt.date.today().isoformat()

def sync_assets(src: Path, dst: Path):
    """Deploy assets from script's assets/ to the public web assets/."""
    if not src.exists():
        return
    dst.mkdir(parents=True, exist_ok=True)
    for p in src.iterdir():
        if p.is_file():
            shutil.copy2(p, dst / p.name)

# =============
# Date handling
# =============
DATE_TOKEN_RE = re.compile(r'(\d{1,2})(?:st|nd|rd|th)?([A-Za-z]+)(\d{4})')
MONTH_ABBR_FOR_EXEC = {1:"Jan",2:"Feb",3:"Mar",4:"Apr",5:"May",6:"Jun",7:"Jul",8:"Aug",9:"Sept",10:"Oct",11:"Nov",12:"Dec"}
MONTH_LONG = {1:"January",2:"February",3:"March",4:"April",5:"May",6:"June",7:"July",8:"August",9:"September",10:"October",11:"November",12:"December"}

def parse_dump_date_token(name: str):
    """
    Parse tokens like '4thSept2025' or '10thSep2025' in a dump folder name and return datetime.date.
    """
    m = DATE_TOKEN_RE.search(name)
    if not m:
        return None
    day = int(m.group(1))
    mon_token = m.group(2).lower()
    year = int(m.group(3))

    table = ["jan","feb","mar","apr","may","jun","jul","aug","sept","oct","nov","dec"]
    month = None
    for key in ("sept","sep","jan","feb","mar","apr","may","jun","jul","aug","oct","nov","dec"):
        if mon_token.startswith(key):
            norm = "sept" if key in ("sept","sep") else key
            month = table.index(norm) + 1
            break
    if not month:
        return None
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None

def ordinal(n: int) -> str:
    return f"{n}{'th' if 11<=n%100<=13 else {1:'st',2:'nd',3:'rd'}.get(n%10,'th')}"

def display_long(d: dt.date) -> str:
    return f"{ordinal(d.day)} {MONTH_LONG[d.month]} {d.year}"

def display_dash(d: dt.date) -> str:
    return f"{d.day:02d}-{d.month:02d}-{d.year}"

def exec_date_token(d: dt.date) -> str:
    """Return token like '7Sept2025' for file naming and links."""
    return f"{d.day}{MONTH_ABBR_FOR_EXEC[d.month]}{d.year}"

# =======================
# Latest dump by NAME date
# =======================
def find_latest_dump_by_name(root: Path) -> Path:
    """
    Choose the latest dump by parsing date tokens in folder names (no mtime).
    If none parse, error out.
    """
    candidates = [p for p in root.iterdir() if p.is_dir()]
    dated = []
    for p in candidates:
        d = parse_dump_date_token(p.name)
        if d:
            dated.append((d, p))
    if not dated:
        sys.exit(f"[ERROR] No parseable dump dates under {root}. Folder names must contain tokens like '4thSept2025'.")
    dated.sort(key=lambda x: x[0], reverse=True)
    return dated[0][1]

# ======
# Email
# ======
def send_mail(cfg: ConfigParser, subj: str, body: str):
    rcpts = [x.strip() for x in cfg.get("mail", "to", fallback="").split(",") if x.strip()]
    if not rcpts:
        return
    method = cfg.get("mail", "method", fallback="mutt").strip().lower()
    try:
        if method == "mutt":
            subprocess.run(["mutt", "-s", subj, *rcpts], input=body.encode(), check=False)
        elif method == "sendmail":
            msg = f"Subject: {subj}\nTo: {', '.join(rcpts)}\n\n{body}"
            subprocess.run(["/usr/sbin/sendmail", "-t", "-oi"], input=msg.encode(), check=False)
        else:
            host = cfg.get("smtp", "host")
            port = cfg.getint("smtp", "port", fallback=587)
            user = cfg.get("smtp", "user", fallback="")
            pwd  = cfg.get("smtp", "password", fallback="")
            sender = cfg.get("smtp", "from", fallback=user or "noreply@example.com")
            use_tls = cfg.getboolean("smtp", "starttls", fallback=True)
            msg = MIMEText(body, "plain", "utf-8")
            msg["Subject"] = subj
            msg["From"] = sender
            msg["To"] = ", ".join(rcpts)
            with smtplib.SMTP(host, port, timeout=45) as s:
                s.ehlo()
                if use_tls:
                    s.starttls()
                    s.ehlo()
                if user:
                    s.login(user, pwd)
                s.sendmail(sender, rcpts, msg.as_string())
    except Exception as e:
        print(f"[WARN] mail failed: {e}", file=sys.stderr)

# =========================
# HTML helpers / placeholders
# =========================
def counts_to_html(counter: Counter) -> str:
    if not counter:
        return '<div class="note">N/A</div>'
    items = []
    for t, n in sorted(counter.items(), key=lambda kv: (-kv[1], kv[0].lower())):
        items.append(f'<li><span class="t">{t}</span><span class="c">{n}</span></li>')
    return '<ul class="typecounts">' + "".join(items) + '</ul>'

def chips_html(items) -> str:
    return ('<div class="chips">' + "".join(f'<span class="chip meta">{x}</span>' for x in items) + '</div>') if items else '<div class="note">No metadata configured.</div>'

def render_template(tpl: str, **kw) -> str:
    out = tpl
    for k, v in kw.items():
        out = out.replace("{{" + k + "}}", str(v))
    return out

def inject_master(tpl: str, reports: list, logo_rel: str, meta_html: str) -> str:
    js = json.dumps(reports, ensure_ascii=False).replace("</script>", "<\\/script>")
    return tpl.replace("{{REPORT_CARDS_JSON}}", js)\
              .replace("{{CADENCE_LOGO}}", logo_rel)\
              .replace("{{METADATA_SIDEBAR_HTML}}", meta_html)

# =========================
# XLSX-only writer utilities
# =========================
def require_openpyxl_or_die(cfg: ConfigParser, logp: Path, context: str):
    try:
        import openpyxl  # noqa: F401
    except Exception as e:
        msg = f"openpyxl not available: {e}\nContext: {context}\nAborting (XLSX-only policy)."
        append_text(logp, "[ERROR] " + msg + "\n")
        send_mail(cfg, "[ProdvsProd] XLSX module missing", msg)
        sys.exit(6)

def write_xlsx_rows(path: Path, sheet_name: str, headers: list, rows: list,
                    cfg: ConfigParser, logp: Path, context: str):
    require_openpyxl_or_die(cfg, logp, context)
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        if headers:
            ws.append(headers)
        for r in rows:
            ws.append(r)
        wb.save(path)
        append_text(logp, f"XLSX written: {path}\n")
    except Exception as e:
        msg = f"Failed to write XLSX at {path}: {e}\nContext: {context}"
        append_text(logp, "[ERROR] " + msg + "\n")
        send_mail(cfg, "[ProdvsProd] XLSX write failed", msg)
        sys.exit(6)

# =========================
# SharePoint link finalizer
# =========================
def finalize_sharepoint_link(pattern: str, exec_token: str) -> str:
    """
    Force .xlsx, drop '(Execution Date)', and fill placeholders.
    """
    if not pattern:
        return ""
    s = pattern
    s = s.replace("{EXEC_DATE}", exec_token)
    s = s.replace("{EXT}", "xlsx")
    s = s.replace("(Execution Date)", "").replace("(execution date)", "")
    s = re.sub(r"\.csv(\b|$)", ".xlsx", s)
    s = re.sub(r"(?<!:)//+", "/", s)
    return s

# =========================
# codecomp.html beautifier
# =========================
from collections import Counter as _Counter
import shutil as _shutil

def beautify_codecomp_html(codecomp_path: Path,
                           out_copy_at_run_root: Path,
                           old_label: str,    # LHS (Previous)
                           new_label: str,    # RHS (Latest)
                           new_counts: _Counter,        # "New" per metadata
                           modified_counts: _Counter):  # "Modified" per metadata
    """
    Overwrite codecomp.html with a minimal, beautiful page:
      - Header with title, LHS/RHS dump names, logo, and back button.
      - One table: [Metadata Type | File Changes (Total) | New | Modified]
        * Metadata Type styled as a blue pill to feel clickable.
      - Footer only.
    """
    metas = sorted(set(new_counts.keys()) | set(modified_counts.keys()), key=lambda s: s.lower())

    def row(meta: str) -> str:
        new_n = int(new_counts.get(meta, 0))
        mod_n = int(modified_counts.get(meta, 0))
        tot_n = new_n + mod_n
        return (
            f"<tr>"
            f'  <td><button class="meta-btn" type="button" aria-label="Open {meta}">{meta}</button></td>'
            f'  <td><span class="num black">{tot_n}</span></td>'
            f'  <td><span class="num green">{new_n}</span></td>'
            f'  <td><span class="num orange">{mod_n}</span></td>'
            f"</tr>"
        )

    rows_html = (
        "\n".join(row(m) for m in metas)
        if metas else '<tr><td colspan="4" class="empty">No differences detected.</td></tr>'
    )

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>Code Comparison Report — {old_label} vs {new_label}</title>
<style>
  :root{{
    --blue:#1d4ed8; --blue-deep:#1e40af;
    --green:#16a34a;
    --orange:#d97706;
    --gray:#6b7280; --ink:#0b0f19; --line:#e6e9ef; --bg:#f8fafc; --white:#fff;
    --radius:16px; --shadow:0 8px 24px rgba(0,0,0,.08),0 2px 6px rgba(0,0,0,.06);
  }}
  *,*::before,*::after{{box-sizing:border-box}}
  body{{margin:0;background:var(--bg);color:var(--ink);font:14px/1.5 system-ui,-apple-system,Segoe UI,Roboto,Ubuntu,Cantarell,Noto Sans,sans-serif}}

  header{{background:linear-gradient(135deg,#0b0f19 0%,var(--blue-deep) 100%);color:#fff}}
  .h-inner{{max-width:1100px;margin:0 auto;padding:14px 18px;display:flex;align-items:center;justify-content:space-between;gap:12px}}
  .h-title{{display:flex;flex-direction:column;gap:6px}}
  .h-title h1{{margin:0;font-size:clamp(18px,2.2vw,24px);font-weight:900;letter-spacing:.2px}}
  .badges{{display:flex;gap:8px;flex-wrap:wrap}}
  .badge{{background:rgba(255,255,255,.14);border:1px solid rgba(255,255,255,.28);color:#fff;padding:4px 10px;border-radius:999px;font-size:12px}}
  .brand{{display:flex;align-items:center;gap:8px}}
  .brand img{{height:32px;display:block;filter:drop-shadow(0 2px 2px rgba(0,0,0,.35))}}
  .back{{display:inline-flex;align-items:center;gap:6px;font-weight:700;font-size:12px;
         background:rgba(255,255,255,.10);border:1px solid rgba(255,255,255,.35);
         border-radius:999px;padding:4px 8px;color:#fff;text-decoration:none}}
  .back:hover{{background:rgba(255,255,255,.18)}}

  main{{max-width:1100px;margin:18px auto;padding:0 18px}}
  .card{{background:#fff;border:1px solid var(--line);border-radius:var(--radius);box-shadow:var(--shadow);padding:18px}}

  table.summary{{width:100%;border-collapse:separate;border-spacing:0 10px}}
  table.summary th, table.summary td{{text-align:left;padding:10px 12px}}
  table.summary thead th{{font-size:12px;color:#111;letter-spacing:.3px;border-bottom:1px solid var(--line)}}
  table.summary tbody tr{{background:#fff;border:1px solid var(--line)}}
  table.summary tbody td{{border-top:1px solid var(--line);border-bottom:1px solid var(--line)}}
  table.summary tbody tr td:first-child{{border-left:1px solid var(--line);border-top-left-radius:12px;border-bottom-left-radius:12px}}
  table.summary tbody tr td:last-child{{border-right:1px solid var(--line);border-top-right-radius:12px;border-bottom-right-radius:12px}}

  .meta-btn{{
    background:var(--blue); color:#fff; border:none; border-radius:999px;
    padding:8px 12px; font-weight:800; cursor:pointer; box-shadow:0 2px 6px rgba(0,0,0,.12);
    transition:transform .15s ease, box-shadow .15s ease, filter .15s ease;
  }}
  .meta-btn:hover{{ transform:translateY(-1px); box-shadow:0 6px 14px rgba(0,0,0,.18); filter:brightness(1.05); }}
  .meta-btn:active{{ transform:translateY(0); box-shadow:0 3px 8px rgba(0,0,0,.14); }}

  .num{{display:inline-block;min-width:30px;text-align:center;font-weight:900}}
  .num.black{{color:#111}}
  .num.green{{color:var(--green)}}
  .num.orange{{color:var(--orange)}}
  .empty{{text-align:center;color:#6b7280}}

  footer{{background:#0b0f19;color:#fff;text-align:center;font-size:12px;padding:6px 10px;margin-top:24px}}
</style>
</head>
<body>
<header>
  <div class="h-inner">
    <div class="h-title">
      <h1>Code Comparison Report</h1>
      <div class="badges">
        <span class="badge">LHS (Previous): <strong>{old_label}</strong></span>
        <span class="badge">RHS (Latest): <strong>{new_label}</strong></span>
      </div>
    </div>
    <div class="brand">
      <img src="../assets/cadence-logo.png" alt="Cadence logo" />
      <a class="back" href="./index.html">← Back</a>
    </div>
  </div>
</header>

<main>
  <section class="card">
    <table class="summary" aria-label="Summary by Metadata Type">
      <thead>
        <tr>
          <th>Metadata Type</th>
          <th>File Changes (Total)</th>
          <th>New</th>
          <th>Modified</th>
        </tr>
      </thead>
      <tbody>
        {rows_html}
      </tbody>
    </table>
  </section>
</main>

<footer>This is the end of report.</footer>
</body>
</html>
"""
    codecomp_path.parent.mkdir(parents=True, exist_ok=True)
    codecomp_path.write_text(html, encoding="utf-8")
    if out_copy_at_run_root:
        out_copy_at_run_root.parent.mkdir(parents=True, exist_ok=True)
        _shutil.copy2(codecomp_path, out_copy_at_run_root)

# ==========================
# No-new-dump early exit
# ==========================
def abort_if_same_dump(old_src: Path, new_src: Path, cfg: ConfigParser, work_path: Path):
    """
    If previous and latest dumps resolve to the same folder, notify and stop.
    - Emails a notice
    - Writes a marker file in work root
    - Prints a message (flushed)
    - Exits with code 7 (non-zero so jobs surface it)
    """
    if old_src.name != new_src.name:
        return
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    msg = (
        "No new production dump found.\n"
        f"Previous and latest both resolve to: {old_src.name}\n"
        f"Root: {old_src.parent}\n"
        "Action: Take a new production dump and rerun."
    )
    send_mail(cfg, "[ProdvsProd] No new dump taken", msg)
    marker = work_path / f"NO_NEW_DUMP_{ts}.txt"
    marker.write_text(msg + "\n", encoding="utf-8")
    print("[NO-OP] " + msg, flush=True)
    sys.exit(7)

# =========
# main flow
# =========
def main():
    here = Path(__file__).resolve().parent
    cfg  = load_conf(here / "setup.conf")

    P = {
        "dumps":  Path(cfg.get("paths", "dumps_root",        fallback="/backup/PROD_DUMP")),
        "work":   Path(cfg.get("paths", "work_root",         fallback="/data/public/ProdvsProdComp")),
        "script": Path(cfg.get("paths", "compare_script_dir",fallback="/software/codedev/misc_script/SFDC/BCompareScript")),
        "last":   Path(cfg.get("paths", "lastdump_file",     fallback=str(here / "lastdumpcomp.date"))),
        "tpl":    Path(cfg.get("paths", "templates_dir",     fallback=str(here / "templates"))),
        "assets": Path(cfg.get("paths", "assets_dir",        fallback=str(here / "assets"))),
    }

    # Deploy assets (logo) to web root
    public_assets = P["work"] / "assets"
    sync_assets(P["assets"], public_assets)

    # Required metadata list
    req_raw = cfg.get("metadata", "current_sfdc_metadata_list_used_in_cadence", fallback="")
    req = [x.strip() for x in req_raw.split(",") if x.strip()]

    # Resolve dumps
    old_name = read_text(P["last"]).strip()
    if not old_name:
        send_mail(cfg, "[ProdvsProd] lastdumpcomp.date is empty", f"Populate {P['last']} with the PREVIOUS dump folder name.")
        sys.exit(1)
    old_src = P["dumps"] / old_name
    if not old_src.exists():
        send_mail(cfg, "[ProdvsProd] previous dump missing", f"Not found: {old_src}")
        sys.exit(1)

    new_src = find_latest_dump_by_name(P["dumps"])  # by name date (not mtime)
    abort_if_same_dump(old_src, new_src, cfg, P["work"])

    # API version
    m = re.search(r"prod_([0-9.]+)_", new_src.name)
    api = m.group(1) if m else cfg.get("ui", "api_version_fallback", fallback="62.0")

    # Title & friendly dates
    def title_from_names(o: str, n: str) -> str:
        d_old = parse_dump_date_token(o)
        d_new = parse_dump_date_token(n)
        if d_old and d_new:
            return f"Production Dump Comparison: {display_long(d_old)} vs. {display_long(d_new)}"
        return f"Production Dump Comparison: {o} vs. {n}"
    title = title_from_names(old_src.name, new_src.name)

    d_old = parse_dump_date_token(old_src.name)
    d_new = parse_dump_date_token(new_src.name)
    prev_human   = display_long(d_old) if d_old else old_src.name
    latest_human = display_long(d_new) if d_new else new_src.name
    prev_dash    = display_dash(d_old) if d_old else ""
    latest_dash  = display_dash(d_new) if d_new else ""

    # Prepare run folder
    day = today_iso()
    run = P["work"] / day
    run.mkdir(parents=True, exist_ok=True)
    log = run / "run.log"
    write_text(log, f"[{dt.datetime.now().isoformat(timespec='seconds')}] Run start\nOld: {old_src}\nNew: {new_src}\n\n")

    # Copy dumps into run folder
    prev_dir   = run / old_src.name
    latest_dir = run / new_src.name
    copy_tree(old_src, prev_dir);   append_text(log, f"Copied -> {prev_dir}\n")
    copy_tree(new_src, latest_dir); append_text(log, f"Copied -> {latest_dir}\n")

    # Parity & required metadata
    tl_prev = top_level_dirs(prev_dir)
    tl_new  = top_level_dirs(latest_dir)
    if tl_prev != tl_new:
        msg = (f"Top-level folder mismatch.\nPrev: {sorted(tl_prev)}\nLatest: {sorted(tl_new)}\nRun: {run}")
        append_text(log, "[ERROR] " + msg + "\n")
        send_mail(cfg, "[ProdvsProd] Folder name mismatch", msg)
        sys.exit(2)
    if req:
        miss_prev = [m for m in req if m not in tl_prev]
        miss_new  = [m for m in req if m not in tl_new]
        if miss_prev or miss_new:
            msg = (f"Required metadata missing.\n"
                   f"Required: {req}\n"
                   f"Missing in PREVIOUS ({old_src.name}): {miss_prev}\n"
                   f"Missing in LATEST ({new_src.name}): {miss_new}\nRun: {run}")
            append_text(log, "[ERROR] " + msg + "\n")
            send_mail(cfg, "[ProdvsProd] Required metadata missing in dumps", msg)
            sys.exit(3)

    # component.list for external compare
    write_text(P["script"] / "component.list", "\n".join(sorted(tl_prev)) + "\n")

    # Prepare compare output dir expected by codecompare script
    comp_dir = run / f"{old_src.name}_vs_{new_src.name}" / "codecomp"
    comp_dir.mkdir(parents=True, exist_ok=True)
    append_text(log, f"Prepared compare output dir: {comp_dir}\n")

    # Optional external compare
    if cfg.getboolean("compare", "run_external_compare", fallback=True):
        sh = P["script"] / "codecompare.posix2.sh"
        if sh.exists():
            cmd = ["bash", str(sh), "-d", "-m", "oo", "-o", old_src.name, "-o", new_src.name]
            append_text(log, f"RUN: {' '.join(cmd)} (cwd={run})\n")
            subprocess.run(cmd, cwd=str(run), check=False)
        else:
            append_text(log, f"[WARN] Missing compare script: {sh}\n")

    # Compute diffs
    prev_files   = {rel: f for rel, f in walk_files(prev_dir)}
    latest_files = {rel: f for rel, f in walk_files(latest_dir)}
    new_rel      = [r for r in latest_files if r not in prev_files]
    removed_rel  = [r for r in prev_files  if r not in latest_files]
    changed_rel  = []
    for r in sorted(set(prev_files) & set(latest_files)):
        a, b = prev_files[r], latest_files[r]
        if a.stat().st_size == b.stat().st_size and sha256(a) == sha256(b):
            continue
        changed_rel.append(r)

    # Per-metadata counts (top-level folder)
    def top_folder(rel: Path) -> str:
        return rel.parts[0] if rel.parts else "(root)"
    cnt_new = Counter(top_folder(r) for r in new_rel)
    cnt_chg = Counter(top_folder(r) for r in changed_rel)
    cnt_rm  = Counter(top_folder(r) for r in removed_rel)

    # --- XLSX reports (NO CSV) ---
    require_openpyxl_or_die(cfg, log, "XLSX-only policy")

    exec_date = dt.date.today()
    token     = exec_date_token(exec_date)  # e.g., 7Sept2025
    base_name = f"ProdDumpComparison_List_{token}"  # no "(Execution Date)"
    full_path = run / (base_name + ".xlsx")

    headers_full = [
        "File List",
        "File Type",                 # "New File" | "Content Mismatch"
        "Dev Team Comment",
        "PS team Comment",
        "DevOps Team Comment",
        "Ticket Number",
        "Ticket Resolution Date",
        "Area (Module)",
        "Details of Change Made",
    ]
    rows_full = (
        [[str(r), "New File", "", "", "", "", "", "", ""] for r in sorted(new_rel)] +
        [[str(r), "Content Mismatch", "", "", "", "", "", "", ""] for r in sorted(changed_rel)]
    )
    write_xlsx_rows(full_path, "Comparison", headers_full, rows_full, cfg, log, "Full comparison report")

    changed_xlsx = run / "changed_files_only.xlsx"
    write_xlsx_rows(changed_xlsx, "Changed Files", ["File List"], [[str(r)] for r in sorted(changed_rel)],
                    cfg, log, "Changed files only")

    new_xlsx = run / "new_files_only.xlsx"
    write_xlsx_rows(new_xlsx, "New Files", ["File List"], [[str(r)] for r in sorted(new_rel)],
                    cfg, log, "New files only")

    removed_xlsx = run / "removed_files_only.xlsx"
    write_xlsx_rows(removed_xlsx, "Removed Files", ["File List"], [[str(r)] for r in sorted(removed_rel)],
                    cfg, log, "Removed files only")

    # Beautify codecomp.html (write regardless of BC output presence)
    bc_out = run / f"{old_src.name}_vs_{new_src.name}" / "codecomp" / "codecomp.html"
    bc_alias = run / "codecomp.html"  # convenience alias for your detail link
    try:
        beautify_codecomp_html(
            codecomp_path=bc_out if bc_out.parent.exists() else bc_alias,
            out_copy_at_run_root=bc_alias,
            old_label=old_src.name,
            new_label=new_src.name,
            new_counts=cnt_new,
            modified_counts=cnt_chg
        )
        append_text(log, "Beautified codecomp.html\n")
    except Exception as e:
        append_text(log, f"[WARN] Beautify codecomp failed: {e}\n")

    # SharePoint link (force XLSX, remove "(Execution Date)")
    sp_pattern = (
        cfg.get("links", "full_report_sharepoint_pattern", fallback="").strip()
        or cfg.get("links", "csv_full_sharepoint_pattern", fallback="").strip()
    )
    full_href = finalize_sharepoint_link(sp_pattern, token) if sp_pattern else f"./{base_name}.xlsx"
    full_label = "Production Modified File Report"

    # Render Detailed page
    dtpl = read_text(P["tpl"] / "detail_template.html")
    if not dtpl:
        sys.exit(f"[ERROR] Missing template: {P['tpl'] / 'detail_template.html'}")

    detail_html = render_template(
        dtpl,
        PAGE_TITLE      = title,
        TITLE_HEADING   = title,
        API_VERSION     = api,
        REPORT_DATE     = dt.date.today().isoformat(),
        PREVIOUS_LABEL  = old_src.name,
        LATEST_LABEL    = new_src.name,
        PREVIOUS_LINK   = f"./{old_src.name}/",
        LATEST_LINK     = f"./{new_src.name}/",
        KPI_NEW         = len(new_rel),
        KPI_CHANGED     = len(changed_rel),
        KPI_REMOVED     = len(removed_rel),
        CSV_FULL        = full_href,
        CSV_FULL_LABEL  = full_label,
        CSV_NEW         = "./new_files_only.xlsx",
        CSV_CHANGED     = "./changed_files_only.xlsx",
        CSV_REMOVED     = "./removed_files_only.xlsx",
        DIFF_HTML_LINK  = "./codecomp.html",
        TYPE_COUNTS_NEW = counts_to_html(cnt_new),
        TYPE_COUNTS_CHANGED = counts_to_html(cnt_chg),
        TYPE_COUNTS_REMOVED = counts_to_html(cnt_rm),
        RUN_LOG_LINK    = "./run.log",
        LOGO_HREF       = "../assets/cadence-logo.png",
        PREVIOUS_DATE_HUMAN = prev_human,
        LATEST_DATE_HUMAN   = latest_human,
        PREVIOUS_DATE_DASH  = prev_dash,
        LATEST_DATE_DASH    = latest_dash,
    )
    write_text(run / "index.html", detail_html)

    # Update master database
    db = P["work"] / "reports.json"
    try:
        existing = json.loads(read_text(db)) if db.exists() else []
    except Exception:
        existing = []

    def pretty_label(n: str) -> str:
        d = parse_dump_date_token(n)
        return display_long(d) if d else n

    entry = {
        "runDate":      day,
        "apiVersion":   api,
        "title":        title,
        "oldDumpLabel": pretty_label(old_src.name),
        "newDumpLabel": pretty_label(new_src.name),
        "detailHref":   f"{day}/index.html",
        "compareHref":  f"{day}/codecomp.html"
    }
    existing.append(entry)
    seen = set(); merged = []
    for r in sorted(existing, key=lambda x: x["runDate"]):
        if r["detailHref"] in seen: continue
        seen.add(r["detailHref"]); merged.append(r)
    write_text(db, json.dumps(merged, ensure_ascii=False, indent=2))

    # Render master index
    mtpl = read_text(P["tpl"] / "index_template.html")
    if not mtpl:
        sys.exit(f"[ERROR] Missing template: {P['tpl'] / 'index_template.html'}")
    master_html = inject_master(mtpl, merged, "assets/cadence-logo.png", chips_html(req))
    write_text(P["work"] / "index.html", master_html)

    # Success email with links (requires public_base_url)
    base_url = cfg.get("links", "public_base_url", fallback="").rstrip("/")
    if base_url:
        master_url  = f"{base_url}/index.html"
        detail_url  = f"{base_url}/{day}/index.html"
        compare_url = f"{base_url}/{day}/codecomp.html"
        send_mail(cfg,
            f"[ProdvsProd] Report built {day}",
            f"Master Index: {master_url}\n"
            f"Detailed Page: {detail_url}\n"
            f"Diff Report:   {compare_url}\n"
        )

    # Update lastdumpcomp.date to NEW dump
    write_text(P["last"], new_src.name + "\n")

    append_text(log, f"[{dt.datetime.now().isoformat(timespec='seconds')}] Run end\n")
    print(f"[DONE] {run}\n  Detailed: {run/'index.html'}\n  Compare:  {run/'codecomp.html'}\n  Master:   {P['work']/'index.html'}")

if __name__ == "__main__":
    main()
